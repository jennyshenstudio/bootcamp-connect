"""Creates cv-word.docx, a sample Word CV for the import tests. Needs: pip install python-docx"""
import os
from docx import Document

d = Document()
d.add_heading('Priya Patel', 0)
d.add_paragraph('Manchester · priya@example.com · github.com/priyapatel')
d.add_heading('Professional Summary', 1)
d.add_paragraph('Business developer with a background in retail analytics. I find customers, test pricing, and build financial models for early-stage products.')
d.add_heading('Work Experience', 1)
d.add_paragraph('Growth Analyst at Northern Retail Digital    Jan 2022 – Present')
d.add_paragraph('Ran pricing experiments that lifted conversion 12%.', style='List Bullet')
d.add_paragraph('Sales Associate, High Street Electronics    06/2019 - 12/2021')
d.add_paragraph('Top seller in electronics for 6 quarters.', style='List Bullet')
d.add_heading('Skills', 1)
d.add_paragraph('Financial modeling, Market research, SQL, Excel, Customer discovery, Pitch decks')
d.add_heading('Projects', 1)
d.add_paragraph('Campus Eats pitch deck')
d.add_paragraph('Pitch and financial model for a student food delivery startup; placed 2nd at a university pitch night.', style='List Bullet')
d.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'cv-word.docx'))
